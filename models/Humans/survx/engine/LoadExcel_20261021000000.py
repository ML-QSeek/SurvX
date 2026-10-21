# -*- coding: utf-8 -*-
"""
LoadExcel.py

中文：从 Init.xlsx 灌库。Matter → SQLite，Energy → DuckDB，_forge → SQLite。
      要求库和表已存在，不建表；没表直接报错，不做兜底。
English: Load Init.xlsx into existing databases. No table creation, no fallback.

版本: 20261021000000
作者: quruyi
时间: 20261021000000

用法:
    python LoadExcel.py <目录> <项目名> --mode append
    python LoadExcel.py <目录> <项目名> --mode replace
"""

import sys
import os
import sqlite3
import duckdb
from openpyxl import load_workbook


# 各 sheet 的字段列表（不含 note）
SHEETS = {
    "Matter": ["key", "value", "version"],
    "_parameter": ["id0", "id1", "value", "idx", "version"],
    "_constraint": ["id0", "id1", "value", "idx", "version"],
    "_relation": ["id0", "id1", "value", "idx", "version"],
    "_sequence": ["_tick", "key", "value"],
    "_experience": ["_from", "_tick", "_type", "_value"],
    "_intent": ["_from", "_tick", "key", "value"],
    "_feedback": ["_from", "_tick", "value"],
    "_language": ["key", "value"],
    "_train": ["_target", "_tick", "key", "value"],
    "_report": ["_type", "_tick", "key", "value"],
    "_llm": ["id", "version", "prompt", "reason"],
    "_register": ["key", "value"],
    "_forge": ["id", "_tick", "_module", "_action", "_target", "_result", "_reason"],
}

# DuckDB 的 Energy 表
DUCK_TABLES = [
    "_parameter", "_constraint", "_relation", "_sequence", "_experience",
    "_intent", "_feedback", "_language", "_train", "_report", "_llm", "_register",
]


def read_sheet(wb, sheet_name, columns):
    """读一个 sheet，返回行列表（每行是字段值列表），跳过 note 列。"""
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    header = [cell.value for cell in ws[1]]
    col_idx = []
    for col in columns:
        if col not in header:
            raise ValueError(f"sheet {sheet_name} 缺少列: {col}")
        col_idx.append(header.index(col))

    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        rows.append([row[i] for i in col_idx])
    return rows


def insert_rows(conn, table, columns, rows, mode, db_kind):
    """插入行。replace 先清空表。"""
    if mode == "replace":
        conn.execute(f"DELETE FROM {table}")

    if not rows:
        return

    cols = ", ".join(columns)
    placeholders = ", ".join(["?"] * len(columns))
    sql = f"INSERT INTO {table} ({cols}) VALUES ({placeholders})"
    for row in rows:
        conn.execute(sql, row)

    if db_kind == "sqlite":
        conn.commit()


def main():
    if len(sys.argv) != 5 or sys.argv[3] != "--mode":
        print("用法: python LoadExcel.py <目录> <项目名> --mode append|replace")
        sys.exit(1)

    base_dir = sys.argv[1]
    project = sys.argv[2]
    mode = sys.argv[4]
    if mode not in ("append", "replace"):
        print(f"未知 mode: {mode}，只支持 append / replace")
        sys.exit(1)

    project_dir = os.path.join(base_dir, project)
    xlsx_path = os.path.join(project_dir, "Init.xlsx")

    if not os.path.exists(xlsx_path):
        print(f"未找到 {xlsx_path}")
        sys.exit(1)

    wb = load_workbook(xlsx_path, data_only=True)

    # ---- SQLite: Matter ----
    sqlite_path = os.path.join(project_dir, f"{project}.db")
    conn = sqlite3.connect(sqlite_path)
    rows = read_sheet(wb, "Matter", SHEETS["Matter"])
    insert_rows(conn, project, SHEETS["Matter"], rows, mode, "sqlite")
    conn.close()
    print(f"Matter: {len(rows)} 行 → {sqlite_path}")

    # ---- SQLite: _forge ----
    forge_path = os.path.join(project_dir, "_forge.db")
    conn = sqlite3.connect(forge_path)
    rows = read_sheet(wb, "_forge", SHEETS["_forge"])
    insert_rows(conn, "_forge", SHEETS["_forge"], rows, mode, "sqlite")
    conn.close()
    print(f"_forge: {len(rows)} 行 → {forge_path}")

    # ---- DuckDB: Energy ----
    duck_path = os.path.join(project_dir, f"_{project}.duckdb")
    duck = duckdb.connect(duck_path)
    for table in DUCK_TABLES:
        rows = read_sheet(wb, table, SHEETS[table])
        insert_rows(duck, table, SHEETS[table], rows, mode, "duckdb")
        print(f"{table}: {len(rows)} 行")
    duck.close()

    print(f"灌库完成，mode={mode}。")


if __name__ == "__main__":
    main()